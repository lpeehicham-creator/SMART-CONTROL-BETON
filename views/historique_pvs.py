import io
import os
import re
import unicodedata
from datetime import datetime, date, timedelta
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
import pandas as pd
import streamlit as st
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
import projets_config


# ==============================================================================
# 1. GESTION UTILISATEURS & SUPABASE
# ==============================================================================
def connecter_utilisateur(supabase, nom_utilisateur, mot_de_passe):
  """Vérifie l'utilisateur et récupère ses droits depuis la table 'users'."""
  try:
    res = (
        supabase.table("users")
        .select("*")
        .eq("username", nom_utilisateur)
        .eq("password", mot_de_passe)
        .execute()
    )
    if res.data:
      user_info = res.data[0]
      st.session_state.update({
          "user_logged": True,
          "user": user_info,
          "username": user_info.get("username"),
          "role": user_info.get("role"),
          "user_role": user_info.get("role"),
          "can_edit": bool(user_info.get("can_edit", False)),
      })
      return True
    st.error("Nom d'utilisateur ou mot de passe incorrect.")
    return False
  except Exception as e:
    st.error(f"Erreur lors de la connexion : {e}")
    return False


def verifier_doublon_num_reception(
    supabase, num_reception, current_beton_id=None
):
  """Vérifie l'existence unique du numéro de réception dans 'suivi_betonnage'."""
  num_clean = str(num_reception or "").strip()
  if not num_clean or num_clean.upper() in ["-", "NONE", "NAN", "N/A"]:
    return False
  try:
    res = (
        projets_config.filtrer_projet_actif(
            supabase.table("suivi_betonnage")
            .select("id, num_reception")
            .eq("num_reception", num_clean)
        ).execute()
    )
    for m in res.data or []:
      if current_beton_id is None or int(m.get("id")) != int(current_beton_id):
        return True
  except Exception as e:
    st.warning(f"Note doublons : {e}")
  return False


def extraire_num_bl(*sources):
  """Extrait le N° BL depuis les sources de données."""
  clefs = {
      "num_bl",
      "bl",
      "num_bon_livraison",
      "n_bl",
      "bon_livraison",
      "num_bl_p",
      "n_bon",
      "bon_de_livraison",
      "code_bl",
  }
  invalid = {"N/A", "NONE", "NAN", "-", ""}
  for src in sources:
    if isinstance(src, dict):
      for k, v in src.items():
        if (k in clefs or "bl" in k.lower() or "bon" in k.lower()) and v:
          v_str = str(v).strip()
          if v_str.upper() not in invalid:
            return v_str
    elif isinstance(src, str):
      match = re.search(r"BL\s*:\s*([^\|]+)", src, re.IGNORECASE)
      if match:
        v_str = match.group(1).strip()
        if v_str.upper() not in invalid:
          return v_str
  return "-"


def calculer_age_jours(date_fab, date_ess, age_defaut=None):
  """Calcule l'âge en jours entre deux dates ou nettoie la valeur existante."""
  try:
    if date_fab and date_ess and str(date_ess).strip().lower() != "en cours":
      df = datetime.strptime(str(date_fab).strip()[:10], "%Y-%m-%d")
      de = datetime.strptime(str(date_ess).strip()[:10], "%Y-%m-%d")
      diff = (de - df).days
      if diff >= 0:
        return diff
  except Exception:
    pass

  if age_defaut is not None:
    val_clean = (
        str(age_defaut)
        .lower()
        .replace("jours", "")
        .replace("jour", "")
        .replace("j", "")
        .strip()
    )
    if val_clean.isdigit():
      return int(val_clean)

  return 28


def est_essai_fendage(item):
  """True si l'éprouvette est un essai de traction par fendage."""
  t = str((item or {}).get("type_essai", "")).strip().lower()
  return "fendage" in t or t.startswith("traction")


def trier_essais_pv(export_data, date_fab):
  """Regroupe les lignes par âge puis par type d'essai (compression avant
  traction), pour que chaque moyenne couvre des lignes contiguës."""
  return sorted(
      export_data,
      key=lambda it: (
          calculer_age_jours(date_fab, it.get("date_essai"), it.get("age")),
          est_essai_fendage(it),
      ),
  )


def nettoyer_nom_fichier(chaine):
  """Remplace les caractères interdits pour les noms de fichiers OS."""
  if not chaine:
    return "PV"
  clean = re.sub(r'[\\/*?:"<>|]', "-", str(chaine).strip())
  return re.sub(r"\s+", "_", clean)


def formater_date_nom_fichier(dt_str):
  """Convertit 'YYYY-MM-DD' en 'DD-MM-YYYY' pour le nom du fichier."""
  if not dt_str or str(dt_str).strip() in ["-", "None", "NaN", "N/A"]:
    return "date_inconnue"
  try:
    dt_obj = datetime.strptime(str(dt_str).strip()[:10], "%Y-%m-%d")
    return dt_obj.strftime("%d-%m-%Y")
  except Exception:
    return str(dt_str).replace("/", "-")


def trouver_logo_lpee():
  """Cherche le logo LPEE (à côté de app.py) et renvoie son chemin, ou None."""
  noms = ["logo.png.jpg", "logo.png", "logo.jpg", "logo.jpeg"]
  ici = os.path.dirname(os.path.abspath(__file__))
  dossiers = [os.getcwd(), ici, os.path.dirname(ici)]
  for dossier in dossiers:
    for nom in noms:
      chemin = os.path.join(dossier, nom)
      if os.path.isfile(chemin):
        return chemin
  return None


# ==============================================================================
# 2. GÉNÉRATION DU PROCÈS-VERBAL PDF (FORMAT LPEE)
# ==============================================================================
@st.cache_data(show_spinner=False)
def generer_pv_excel(export_data, infos_header):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "PV"
  ws.sheet_view.showGridLines = False
  ws.page_setup.orientation = "portrait"
  ws.page_setup.paperSize = ws.PAPERSIZE_A4
  ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
  ws.sheet_properties.pageSetUpPr.fitToPage = True
  ws.page_setup.fitToWidth = 1
  ws.page_setup.fitToHeight = 1

  DARK_FILL = PatternFill("solid", fgColor="1F4E78")
  TABLE_FILL = PatternFill("solid", fgColor="D9E1F2")
  LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
  BLACK = "000000"
  THIN = Side(style="thin", color="808080")
  BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

  default_bl = extraire_num_bl(infos_header)

  def clean_na(val, fallback=default_bl):
    v = str(val).strip() if val is not None else ""
    return fallback if v.upper() in ["N/A", "NONE", "NAN", "", "-"] else val

  def set_cell(row, col, value, bold=False, size=8.5, fill=None, color=BLACK,
               align="center", wrap=False, italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = BORDER_ALL
    if fill is not None:
      c.fill = fill
    return c

  def merge(r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

  widths = {"A": 16, "B": 12, "C": 12, "D": 10, "E": 18, "F": 14, "G": 12, "H": 12}
  for col, w in widths.items():
    ws.column_dimensions[col].width = w

  merge(1, 1, 1, 4)
  set_cell(1, 1, "LPEE / CTR CSB", bold=True, size=11, fill=DARK_FILL, color="FFFFFF")
  set_cell(1, 5, "RE N° :", bold=True)
  merge(1, 6, 1, 7)
  set_cell(1, 6, clean_na(infos_header.get("re_num"), "25/260/LGV/ B/"), bold=True, align="right")
  ref_h1 = clean_na(
      infos_header.get("num_reception")
      or infos_header.get("ref_controle")
      or infos_header.get("reference"),
      "B/406",
  )
  set_cell(1, 8, ref_h1, bold=True, align="left")

  merge(2, 1, 3, 4)
  set_cell(2, 1, "Laboratoire de Contrôle Externe", bold=True, fill=DARK_FILL, color="FFFFFF")
  set_cell(2, 5, "DOSSIER :", bold=True)
  merge(2, 6, 2, 8)
  set_cell(2, 6, clean_na(infos_header.get("dossier"), projets_config.dossier_pv("2025-260-05985-2025-0247")))
  set_cell(3, 5, "CLIENT :", bold=True)
  merge(3, 6, 3, 8)
  set_cell(3, 6, clean_na(infos_header.get("client"), projets_config.client_projet()), bold=True)

  merge(4, 1, 4, 8)
  set_cell(4, 1, "ESSAIS MECANIQUES SUR BETON HYDRAULIQUE", bold=True, size=13, fill=DARK_FILL, color="FFFFFF")

  est_fendage = any(est_essai_fendage(item) for item in export_data)
  est_compression = (not export_data) or any(not est_essai_fendage(item) for item in export_data)
  merge(5, 1, 5, 4)
  set_cell(5, 1, f"[{'X' if est_compression else ' '}] COMPRESSION NF EN 12390-3 (2019)", bold=True)
  merge(5, 5, 5, 8)
  set_cell(5, 5, f"[{'X' if est_fendage else ' '}] TRACTION PAR FENDAGE NF EN 12390-6 (2019)", bold=True)

  merge(6, 1, 6, 6)
  set_cell(6, 1, "Presse : Marque: Controls", bold=True, align="right")
  merge(6, 7, 6, 8)
  set_cell(6, 7, "Classe : A", bold=True)

  date_fab_header = clean_na(infos_header.get("date_coulee"), "-")
  set_cell(7, 1, "Date de\nprélèvement", bold=True, fill=LABEL_FILL, wrap=True)
  set_cell(7, 2, str(date_fab_header), bold=True)
  merge(7, 3, 7, 4)
  set_cell(7, 3, "Lieu de\nprélèvement", bold=True, fill=LABEL_FILL, wrap=True)
  merge(7, 5, 7, 8)
  set_cell(7, 5, clean_na(infos_header.get("lieu_prelevement", infos_header.get("ouvrage")), "-"), bold=True)
  ws.row_dimensions[7].height = 30

  set_cell(8, 1, "Chantier", bold=True, fill=LABEL_FILL)
  merge(8, 2, 8, 4)
  set_cell(8, 2, clean_na(
      infos_header.get("chantier"),
      projets_config.chantier_pv("LGV-Travaux d'exécution de terrassement, ouvrages d'art et rétablissement de communication entre PK 5+500 et PK 10+000-GARE CASA SUD."),
  ), size=7, wrap=True)
  merge(8, 5, 8, 6)
  set_cell(8, 5, "Type de béton", bold=True, fill=LABEL_FILL)
  merge(8, 7, 8, 8)
  set_cell(8, 7, str(clean_na(infos_header.get("classe_beton"), "C35/45")).upper(), bold=True)
  ws.row_dimensions[8].height = 34

  merge(9, 1, 9, 2)
  set_cell(9, 1, clean_na(infos_header.get("centrale"), "Centrale à Béton"), bold=True, fill=LABEL_FILL)
  set_cell(9, 3, "- Dimensions", align="left")
  merge(9, 4, 9, 8)
  set_cell(9, 4, clean_na(infos_header.get("forme"), "Cylindrique 150x300"), bold=True)

  merge(10, 1, 10, 2)
  set_cell(10, 1, "Affaissement au cône d'abrams NF EN 12350-2", size=7, fill=LABEL_FILL, wrap=True)
  set_cell(10, 3, str(clean_na(infos_header.get("affaissement"), "-")), bold=True)
  set_cell(10, 4, "- Mode confection", align="left", size=7)
  merge(10, 5, 10, 8)
  set_cell(10, 5, "Par vibration NF EN 12390-2 (2019)", bold=True)
  ws.row_dimensions[10].height = 26

  merge(11, 1, 11, 2)
  set_cell(11, 1, "Température °C", bold=True, fill=LABEL_FILL)
  set_cell(11, 3, str(clean_na(infos_header.get("temperature"), "-")), bold=True)
  set_cell(11, 4, "- Mode conservation", align="left", size=7)
  merge(11, 5, 11, 8)
  set_cell(11, 5, "au laboratoire par immersion dans l'eau NF EN 12390-2 (2019) à 20°C ± 2°C", bold=True, size=7.5, wrap=True)
  ws.row_dimensions[11].height = 26

  tech = clean_na(
      infos_header.get("technicien_prelevement")
      or infos_header.get("preleve_par")
      or infos_header.get("technicien"),
      "Technicien LPEE",
  )
  merge(12, 1, 12, 3)
  set_cell(12, 1, f"prélèvement effectué par {tech}", size=7, fill=LABEL_FILL, wrap=True)
  merge(12, 4, 12, 5)
  set_cell(12, 4, "N° de bon de livraison", bold=True, fill=LABEL_FILL)
  merge(12, 6, 12, 8)
  set_cell(12, 6, default_bl, bold=True)

  merge(13, 1, 14, 1)
  set_cell(13, 1, "Réf,", bold=True, fill=TABLE_FILL)
  merge(13, 2, 13, 3)
  set_cell(13, 2, "Date", bold=True, fill=TABLE_FILL)
  set_cell(14, 2, "Fabri", bold=True, fill=TABLE_FILL)
  set_cell(14, 3, "Essai", bold=True, fill=TABLE_FILL)
  merge(13, 4, 14, 4)
  set_cell(13, 4, "Age (jours)", bold=True, fill=TABLE_FILL)
  merge(13, 5, 14, 5)
  set_cell(13, 5, "Charge rupture(KN)", bold=True, fill=TABLE_FILL)
  merge(13, 6, 13, 8)
  set_cell(13, 6, "Résistance (MPa)", bold=True, fill=TABLE_FILL)
  set_cell(14, 6, "Compression", bold=True, fill=TABLE_FILL)
  set_cell(14, 7, "Traction", bold=True, fill=TABLE_FILL)
  set_cell(14, 8, "Moyenne", bold=True, fill=TABLE_FILL)

  ligne_courante = 15
  groupes_lots = {}
  export_data = trier_essais_pv(export_data, date_fab_header)
  for item in export_data:
    fend = est_essai_fendage(item)
    f_kn = float(item.get("force_kn", 0.0) or 0.0)
    is_en_cours = str(item.get("statut", "")).lower() == "en cours" or f_kn == 0.0
    dt_essai = item.get("date_essai")
    age_val = calculer_age_jours(date_fab_header, dt_essai, item.get("age"))

    date_essai_affichage = "-"
    if not is_en_cours and dt_essai and str(dt_essai).strip() not in ["-", "", "None", "NaN"]:
      date_essai_affichage = str(clean_na(dt_essai, "-"))
    else:
      try:
        df_obj = datetime.strptime(str(date_fab_header).strip()[:10], "%Y-%m-%d")
        date_essai_affichage = (df_obj + timedelta(days=int(age_val))).strftime("%Y-%m-%d")
      except Exception:
        date_essai_affichage = "-"

    set_cell(ligne_courante, 1, str(item.get("repere_eprouvette", "B/01")))
    set_cell(ligne_courante, 2, str(date_fab_header))
    set_cell(ligne_courante, 3, date_essai_affichage)
    set_cell(ligne_courante, 4, str(age_val))
    col_val, col_autre = (7, 6) if fend else (6, 7)
    if is_en_cours:
      set_cell(ligne_courante, 5, "En cours")
      set_cell(ligne_courante, col_val, "En cours")
    else:
      set_cell(ligne_courante, 5, f"{f_kn:.1f}")
      set_cell(ligne_courante, col_val, f"{float(item.get('fc_mpa', 0.0) or 0.0):.1f}")
    set_cell(ligne_courante, col_autre, "-")

    cle = f"{age_val}_{dt_essai}_{'T' if fend else 'C'}"
    groupes_lots.setdefault(cle, {"lignes": [], "en_cours": is_en_cours, "age": age_val, "fend": fend})["lignes"].append(ligne_courante)
    ligne_courante += 1

  a_des_28j, moyenne_28j_val, est_en_cours_28j = False, None, False
  for gdata in groupes_lots.values():
    lignes, age = sorted(gdata["lignes"]), gdata["age"]
    start_r, end_r = min(lignes), max(lignes)
    contigu = (end_r - start_r + 1) == len(lignes)

    if gdata["en_cours"]:
      valeur_moy = "En cours"
    else:
      vals = []
      for li in lignes:
        try:
          vals.append(float(ws.cell(row=li, column=7 if gdata["fend"] else 6).value))
        except (ValueError, TypeError):
          pass
      moy = round(sum(vals) / len(vals), 1) if vals else 0.0
      valeur_moy = f"{moy:.1f}"
      if int(age) >= 28 and not gdata["fend"]:
        moyenne_28j_val = moy

    if contigu and start_r != end_r:
      merge(start_r, 8, end_r, 8)
      set_cell(start_r, 8, valeur_moy, bold=True)
    else:
      for li in lignes:
        set_cell(li, 8, valeur_moy, bold=True)

    if int(age) >= 28 and not gdata["fend"]:
      a_des_28j = True
      if gdata["en_cours"]:
        est_en_cours_28j = True

  seuil = next(
      (s for k, s in [("C25/30", 25.0), ("C30/37", 30.0), ("C35/45", 35.0), ("C40/50", 40.0)]
       if k in str(clean_na(infos_header.get("classe_beton"), "C35/45")).upper()),
      35.0,
  )
  if not a_des_28j or est_en_cours_28j or moyenne_28j_val is None:
    comment_valeur = "PERFORMANCES MECANIQUES A 28 JOURS SERONT DONNES ULTERIEUREMENT."
  elif moyenne_28j_val >= seuil:
    comment_valeur = "PERFORMANCES MECANIQUES A 28 JOURS SONT CONFORMES"
  else:
    comment_valeur = "PERFORMANCES MECANIQUES NON CONFORMES"

  row_comment = ligne_courante
  set_cell(row_comment, 1, "Commentaire :", bold=True, fill=LABEL_FILL, align="left")
  merge(row_comment, 2, row_comment, 8)
  set_cell(row_comment, 2, comment_valeur, bold=True, align="left")

  # Cases de visa : signataires propres au chantier (créés avec le chantier)
  visas_pv = projets_config.visas_pv()
  cols_visa = projets_config.colonnes_visas(len(visas_pv))
  row_visa_titre = row_comment + 1
  row_visa_nom = row_visa_titre + 1
  for (titre_v, nom_v), (c1_v, c2_v) in zip(visas_pv, cols_visa):
    merge(row_visa_titre, c1_v, row_visa_titre, c2_v)
    set_cell(row_visa_titre, c1_v, titre_v, bold=True)
    merge(row_visa_nom, c1_v, row_visa_nom, c2_v)
    set_cell(row_visa_nom, c1_v, nom_v, bold=True, align="center")
  ws.row_dimensions[row_visa_nom].height = 60

  logo_path = trouver_logo_lpee()
  if logo_path:
    try:
      from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
      from openpyxl.drawing.xdr import XDRPositiveSize2D
      from openpyxl.utils.units import pixels_to_EMU

      # Bloc "Laboratoire de Contrôle Externe" = A2:D3 : logo à gauche,
      # texte à droite du logo.
      ws.row_dimensions[2].height = 30
      ws.row_dimensions[3].height = 30
      ws["A2"].alignment = Alignment(
          horizontal="right", vertical="center", wrap_text=True, indent=2
      )
      larg_bloc_px = int(sum(widths[c] for c in "ABCD") * 7 + 5 * 4)
      haut_bloc_px = int((30 + 30) * 96 / 72)
      img = XLImage(logo_path)
      ratio = img.width / img.height if img.height else 1
      h_px = haut_bloc_px - 12
      w_px = int(h_px * ratio)
      if w_px > larg_bloc_px * 0.4:
        w_px = int(larg_bloc_px * 0.4)
        h_px = int(w_px / ratio)
      img.anchor = OneCellAnchor(
          _from=AnchorMarker(
              col=0, row=1,
              colOff=pixels_to_EMU(6),
              rowOff=pixels_to_EMU(max((haut_bloc_px - h_px) // 2, 0)),
          ),
          ext=XDRPositiveSize2D(pixels_to_EMU(w_px), pixels_to_EMU(h_px)),
      )
      ws.add_image(img)
    except Exception:
      pass  # le PV reste généré même si le logo pose problème

  ws.print_area = f"A1:H{row_visa_nom}"

  buf = io.BytesIO()
  wb.save(buf)
  buf.seek(0)
  return buf.getvalue()


def generer_pv_pdf(export_data, infos_header):
  buf = io.BytesIO()
  left_m = right_m = 0.3 * inch
  top_m = bottom_m = 0.4 * inch

  doc = SimpleDocTemplate(
      buf,
      pagesize=A4,
      leftMargin=left_m,
      rightMargin=right_m,
      topMargin=top_m,
      bottomMargin=bottom_m,
  )

  page_width = A4[0] - left_m - right_m
  base_widths = [16, 12, 12, 10, 18, 14, 12, 12]
  total_units = sum(base_widths)
  col_widths = [page_width * (w / total_units) for w in base_widths]

  DARK = colors.HexColor("#1F4E78")
  TABLE_BG = colors.HexColor("#D9E1F2")
  LABEL_BG = colors.HexColor("#F2F2F2")
  WHITE = colors.white
  BLACK = colors.black

  def P(text, size=7.5, bold=False, align="CENTER", color=BLACK):
    align_map = {"CENTER": TA_CENTER, "LEFT": TA_LEFT, "RIGHT": TA_RIGHT}
    style = ParagraphStyle(
        name="cell",
        fontName="Helvetica-Bold" if bold else "Helvetica",
        fontSize=size,
        leading=size * 1.2,
        alignment=align_map.get(align, TA_CENTER),
        textColor=color,
    )
    return Paragraph(str(text), style)

  default_bl = extraire_num_bl(infos_header)

  def clean_na(val, fallback=default_bl):
    v = str(val).strip() if val is not None else ""
    return fallback if v.upper() in ["N/A", "NONE", "NAN", "", "-"] else val

  def blank_row():
    return ["" for _ in range(8)]

  data = []
  spans, bg, fonts, aligns, valigns = [], [], [], [], []

  r = blank_row()
  r[0] = "LPEE / CTR CSB"
  r[4] = "RE N° :"
  r[5] = clean_na(infos_header.get("re_num"), "25/260/LGV/ B/")
  ref_h1 = clean_na(
      infos_header.get("num_reception")
      or infos_header.get("ref_controle")
      or infos_header.get("reference"),
      "B/406",
  )
  r[7] = ref_h1
  data.append(r)
  row0 = len(data) - 1
  spans += [(0, row0, 3, row0), (5, row0, 6, row0)]
  bg.append((0, row0, 3, row0, DARK))
  fonts.append((0, row0, 3, row0, "Helvetica-Bold", 9, WHITE))
  fonts.append((4, row0, 4, row0, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((5, row0, 6, row0, "RIGHT"))
  fonts.append((7, row0, 7, row0, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((7, row0, 7, row0, "LEFT"))

  r = blank_row()
  r[0] = "Laboratoire de Contrôle Externe"
  r[4] = "DOSSIER :"
  r[5] = clean_na(infos_header.get("dossier"), projets_config.dossier_pv("2025-260-05985-2025-0247"))
  data.append(r)
  row1 = len(data) - 1

  r = blank_row()
  r[4] = "CLIENT :"
  r[5] = clean_na(infos_header.get("client"), projets_config.client_projet())
  data.append(r)
  row2 = len(data) - 1

  spans.append((0, row1, 3, row2))
  bg.append((0, row1, 3, row2, DARK))
  fonts.append((0, row1, 3, row2, "Helvetica-Bold", 9, WHITE))
  spans += [(5, row1, 7, row1), (5, row2, 7, row2)]
  fonts.append((4, row1, 4, row1, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((4, row2, 4, row2, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((5, row1, 7, row1, "Helvetica", 8.5, BLACK))
  fonts.append((5, row2, 7, row2, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  r[0] = "ESSAIS MECANIQUES SUR BETON HYDRAULIQUE"
  data.append(r)
  row3 = len(data) - 1
  spans.append((0, row3, 7, row3))
  bg.append((0, row3, 7, row3, DARK))
  fonts.append((0, row3, 7, row3, "Helvetica-Bold", 11, WHITE))

  est_fendage = any(est_essai_fendage(item) for item in export_data)
  est_compression = (not export_data) or any(not est_essai_fendage(item) for item in export_data)
  r = blank_row()
  r[0] = f"[{'X' if est_compression else ' '}] COMPRESSION NF EN 12390-3 (2019)"
  r[4] = f"[{'X' if est_fendage else ' '}] TRACTION PAR FENDAGE NF EN 12390-6 (2019)"
  data.append(r)
  row4 = len(data) - 1
  spans += [(0, row4, 3, row4), (4, row4, 7, row4)]
  fonts.append((0, row4, 7, row4, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  r[0] = "Presse : Marque: Controls"
  r[6] = "Classe : A"
  data.append(r)
  row5 = len(data) - 1
  spans += [(0, row5, 5, row5), (6, row5, 7, row5)]
  fonts.append((0, row5, 7, row5, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((0, row5, 5, row5, "RIGHT"))

  date_fab_header = clean_na(infos_header.get("date_coulee"), "-")
  r = blank_row()
  r[0] = "Date de\nprélèvement"
  r[1] = str(date_fab_header)
  r[2] = "Lieu de\nprélèvement"
  r[4] = P(
      clean_na(
          infos_header.get("lieu_prelevement", infos_header.get("ouvrage")), "-"
      ),
      size=8.5,
      bold=True,
  )
  data.append(r)
  row6 = len(data) - 1
  spans += [(2, row6, 3, row6), (4, row6, 7, row6)]
  bg += [(0, row6, 0, row6, LABEL_BG), (2, row6, 3, row6, LABEL_BG)]
  fonts.append((0, row6, 0, row6, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((1, row6, 1, row6, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((2, row6, 3, row6, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((4, row6, 7, row6, "Helvetica", 8.5, BLACK))

  r = blank_row()
  r[0] = "Chantier"
  r[1] = P(
      clean_na(
          infos_header.get("chantier"),
          projets_config.chantier_pv("LGV-Travaux d'exécution de terrassement, ouvrages d'art et rétablissement de communication entre PK 5+500 et PK 10+000-GARE CASA SUD."),
      ),
      size=7,
  )
  r[4] = "Type de béton"
  r[6] = str(clean_na(infos_header.get("classe_beton"), "C35/45")).upper()
  data.append(r)
  row7 = len(data) - 1
  spans += [(1, row7, 3, row7), (4, row7, 5, row7), (6, row7, 7, row7)]
  bg += [(0, row7, 0, row7, LABEL_BG), (4, row7, 5, row7, LABEL_BG)]
  fonts.append((0, row7, 0, row7, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((1, row7, 3, row7, "Helvetica", 7.5, BLACK))
  fonts.append((4, row7, 5, row7, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((6, row7, 7, row7, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  r[0] = clean_na(infos_header.get("centrale"), "Centrale à Béton")
  r[2] = "- Dimensions"
  r[3] = clean_na(infos_header.get("forme"), "Cylindrique 150x300")
  data.append(r)
  row8 = len(data) - 1
  spans += [(0, row8, 1, row8), (3, row8, 7, row8)]
  bg.append((0, row8, 1, row8, LABEL_BG))
  fonts.append((0, row8, 1, row8, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((2, row8, 2, row8, "LEFT"))
  fonts.append((3, row8, 7, row8, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  r[0] = P("Affaissement au cône d'abrams NF EN 12350-2", size=7)
  r[2] = str(clean_na(infos_header.get("affaissement"), "-"))
  r[3] = P("- Mode confection", size=7, align="LEFT")
  r[4] = "Par vibration NF EN 12390-2 (2019)"
  data.append(r)
  row9 = len(data) - 1
  spans += [(0, row9, 1, row9), (4, row9, 7, row9)]
  bg.append((0, row9, 1, row9, LABEL_BG))
  fonts.append((0, row9, 1, row9, "Helvetica", 7.5, BLACK))
  fonts.append((2, row9, 2, row9, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((3, row9, 3, row9, "LEFT"))
  fonts.append((4, row9, 7, row9, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  r[0] = "Température °C"
  r[2] = str(clean_na(infos_header.get("temperature"), "-"))
  r[3] = P("- Mode conservation", size=7, align="LEFT")
  r[4] = P(
      "au laboratoire par immersion dans l'eau NF EN 12390-2 (2019) à 20°C"
      " ± 2°C",
      size=7.5,
      bold=True,
  )
  data.append(r)
  row10 = len(data) - 1
  spans += [(0, row10, 1, row10), (4, row10, 7, row10)]
  bg.append((0, row10, 1, row10, LABEL_BG))
  fonts.append((0, row10, 1, row10, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((2, row10, 2, row10, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((3, row10, 3, row10, "LEFT"))
  fonts.append((4, row10, 7, row10, "Helvetica-Bold", 7.5, BLACK))

  tech = clean_na(
      infos_header.get("technicien_prelevement")
      or infos_header.get("preleve_par")
      or infos_header.get("technicien"),
      "Technicien LPEE",
  )
  r = blank_row()
  r[0] = P(f"prélèvement effectué par {tech}", size=7)
  r[3] = "N° de bon de livraison"
  r[5] = default_bl
  data.append(r)
  row11 = len(data) - 1
  spans += [(0, row11, 2, row11), (3, row11, 4, row11), (5, row11, 7, row11)]
  bg += [(0, row11, 2, row11, LABEL_BG), (3, row11, 4, row11, LABEL_BG)]
  fonts.append((0, row11, 2, row11, "Helvetica", 7.5, BLACK))
  fonts.append((3, row11, 4, row11, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((5, row11, 7, row11, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  r[0] = "Réf,"
  r[1] = "Date"
  r[3] = "Age (jours)"
  r[4] = "Charge rupture(KN)"
  r[5] = "Résistance (MPa)"
  data.append(r)
  row12 = len(data) - 1

  r = blank_row()
  r[1] = "Fabri"
  r[2] = "Essai"
  r[5] = "Compression"
  r[6] = "Traction"
  r[7] = "Moyenne"
  data.append(r)
  row13 = len(data) - 1

  spans += [
      (0, row12, 0, row13),
      (1, row12, 2, row12),
      (3, row12, 3, row13),
      (4, row12, 4, row13),
      (5, row12, 7, row12),
  ]
  bg.append((0, row12, 7, row13, TABLE_BG))
  fonts.append((0, row12, 7, row13, "Helvetica-Bold", 8.5, BLACK))

  row_indices_body = []
  groupes_lots = {}
  export_data = trier_essais_pv(export_data, date_fab_header)
  for item in export_data:
    fend = est_essai_fendage(item)
    f_kn = float(item.get("force_kn", 0.0) or 0.0)
    is_en_cours = (
        str(item.get("statut", "")).lower() == "en cours" or f_kn == 0.0
    )
    dt_essai = item.get("date_essai")
    age_val = calculer_age_jours(date_fab_header, dt_essai, item.get("age"))

    date_essai_affichage = "-"
    if (
        not is_en_cours
        and dt_essai
        and str(dt_essai).strip() not in ["-", "", "None", "NaN"]
    ):
      date_essai_affichage = str(clean_na(dt_essai, "-"))
    else:
      try:
        df_obj = datetime.strptime(
            str(date_fab_header).strip()[:10], "%Y-%m-%d"
        )
        date_essai_affichage = (
            df_obj + timedelta(days=int(age_val))
        ).strftime("%Y-%m-%d")
      except Exception:
        date_essai_affichage = "-"

    r = blank_row()
    r[0] = str(item.get("repere_eprouvette", "B/01"))
    r[1] = str(date_fab_header)
    r[2] = date_essai_affichage
    r[3] = str(age_val)
    i_val, i_autre = (6, 5) if fend else (5, 6)
    if is_en_cours:
      r[4] = "En cours"
      r[i_val] = "En cours"
    else:
      r[4] = f"{f_kn:.1f}"
      r[i_val] = f"{float(item.get('fc_mpa', 0.0) or 0.0):.1f}"
    r[i_autre] = "-"
    data.append(r)
    r_idx = len(data) - 1
    row_indices_body.append(r_idx)
    fonts.append((0, r_idx, 7, r_idx, "Helvetica", 8.5, BLACK))

    cle = f"{age_val}_{dt_essai}_{'T' if fend else 'C'}"
    groupes_lots.setdefault(
        cle, {"lignes": [], "en_cours": is_en_cours, "age": age_val, "fend": fend}
    )["lignes"].append(r_idx)

  a_des_28j, moyenne_28j_val, est_en_cours_28j = False, None, False
  for gdata in groupes_lots.values():
    lignes, age = sorted(gdata["lignes"]), gdata["age"]
    start_r, end_r = min(lignes), max(lignes)
    contigu = (end_r - start_r + 1) == len(lignes)

    if gdata["en_cours"]:
      valeur_moy = "En cours"
    else:
      vals = []
      for li in lignes:
        try:
          vals.append(float(data[li][6 if gdata["fend"] else 5]))
        except (ValueError, TypeError):
          pass
      moy = round(sum(vals) / len(vals), 1) if vals else 0.0
      valeur_moy = f"{moy:.1f}"
      if int(age) >= 28 and not gdata["fend"]:
        moyenne_28j_val = moy

    if contigu and start_r != end_r:
      spans.append((7, start_r, 7, end_r))
      data[start_r][7] = valeur_moy
      fonts.append((7, start_r, 7, end_r, "Helvetica-Bold", 8.5, BLACK))
    else:
      for li in lignes:
        data[li][7] = valeur_moy
        fonts.append((7, li, 7, li, "Helvetica-Bold", 8.5, BLACK))

    if int(age) >= 28 and not gdata["fend"]:
      a_des_28j = True
      if gdata["en_cours"]:
        est_en_cours_28j = True

  seuil = next(
      (
          s
          for k, s in [
              ("C25/30", 25.0),
              ("C30/37", 30.0),
              ("C35/45", 35.0),
              ("C40/50", 40.0),
          ]
          if k
          in str(clean_na(infos_header.get("classe_beton"), "C35/45")).upper()
      ),
      35.0,
  )
  if not a_des_28j or est_en_cours_28j or moyenne_28j_val is None:
    comment_valeur = (
        "PERFORMANCES MECANIQUES A 28 JOURS SERONT DONNES ULTERIEUREMENT."
    )
  elif moyenne_28j_val >= seuil:
    comment_valeur = "PERFORMANCES MECANIQUES A 28 JOURS SONT CONFORMES"
  else:
    comment_valeur = "PERFORMANCES MECANIQUES NON CONFORMES"

  r = blank_row()
  r[0] = "Commentaire :"
  r[1] = P(comment_valeur, size=8.5, bold=True, align="LEFT")
  data.append(r)
  row_comment = len(data) - 1
  spans.append((1, row_comment, 7, row_comment))
  bg.append((0, row_comment, 0, row_comment, LABEL_BG))
  fonts.append((0, row_comment, 0, row_comment, "Helvetica-Bold", 8.5, BLACK))
  fonts.append((1, row_comment, 7, row_comment, "Helvetica-Bold", 8.5, BLACK))
  aligns.append((0, row_comment, 0, row_comment, "LEFT"))
  aligns.append((1, row_comment, 7, row_comment, "LEFT"))

  # Cases de visa : signataires propres au chantier (créés avec le chantier)
  visas_pv = projets_config.visas_pv()
  cols_visa = [(a_ - 1, b_ - 1) for a_, b_ in projets_config.colonnes_visas(len(visas_pv))]

  r = blank_row()
  for (titre_v, _nom_v), (c1_v, _c2_v) in zip(visas_pv, cols_visa):
    r[c1_v] = titre_v
  data.append(r)
  row_visa_titre = len(data) - 1
  for (c1_v, c2_v) in cols_visa:
    spans.append((c1_v, row_visa_titre, c2_v, row_visa_titre))
    fonts.append((c1_v, row_visa_titre, c2_v, row_visa_titre, "Helvetica-Bold", 8.5, BLACK))

  r = blank_row()
  for (_titre_v, nom_v), (c1_v, _c2_v) in zip(visas_pv, cols_visa):
    r[c1_v] = nom_v
  data.append(r)
  row_visa_nom = len(data) - 1
  for (c1_v, c2_v) in cols_visa:
    spans.append((c1_v, row_visa_nom, c2_v, row_visa_nom))
    fonts.append((c1_v, row_visa_nom, c2_v, row_visa_nom, "Helvetica-Bold", 9, BLACK))
    valigns.append((c1_v, row_visa_nom, c2_v, row_visa_nom, "TOP"))

  rows_auto_hauteur = {row6, row7, row9, row10, row11, row_comment}
  row_heights = []
  for i in range(len(data)):
    if i in rows_auto_hauteur:
      row_heights.append(None)
    elif i == row_visa_nom:
      row_heights.append(48)
    else:
      row_heights.append(16)

  style_cmds = [
      ("GRID", (0, 0), (-1, -1), 0.5, BLACK),
      ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
      ("ALIGN", (0, 0), (-1, -1), "CENTER"),
      ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
      ("FONTSIZE", (0, 0), (-1, -1), 8),
      ("TOPPADDING", (0, 0), (-1, -1), 2),
      ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
      ("LEFTPADDING", (0, 0), (-1, -1), 2),
      ("RIGHTPADDING", (0, 0), (-1, -1), 2),
  ]
  for (c1, r1, c2, r2) in spans:
    style_cmds.append(("SPAN", (c1, r1), (c2, r2)))
  for (c1, r1, c2, r2, color) in bg:
    style_cmds.append(("BACKGROUND", (c1, r1), (c2, r2), color))
  for (c1, r1, c2, r2, fname, fsize, fcolor) in fonts:
    style_cmds.append(("FONTNAME", (c1, r1), (c2, r2), fname))
    style_cmds.append(("FONTSIZE", (c1, r1), (c2, r2), fsize))
    style_cmds.append(("TEXTCOLOR", (c1, r1), (c2, r2), fcolor))
  for (c1, r1, c2, r2, al) in aligns:
    style_cmds.append(("ALIGN", (c1, r1), (c2, r2), al))
  for (c1, r1, c2, r2, va) in valigns:
    style_cmds.append(("VALIGN", (c1, r1), (c2, r2), va))
  table_style = TableStyle(style_cmds)

  page_height_dispo = A4[1] - top_m - bottom_m

  table1 = Table(data, colWidths=col_widths, rowHeights=row_heights)
  table1.setStyle(table_style)
  _, hauteur_naturelle = table1.wrap(page_width, page_height_dispo * 10)

  if hauteur_naturelle > 0 and hauteur_naturelle < page_height_dispo:
    hauteurs_reelles = list(table1._rowHeights)
    facteur = min((page_height_dispo * 0.97) / hauteur_naturelle, 3.5)
    row_heights_final = [h * facteur for h in hauteurs_reelles]

    table_verif = Table(data, colWidths=col_widths, rowHeights=row_heights_final)
    table_verif.setStyle(table_style)
    _, hauteur_finale = table_verif.wrap(page_width, page_height_dispo * 10)
    if hauteur_finale > page_height_dispo:
      correction = (page_height_dispo * 0.97) / hauteur_finale
      row_heights_final = [h * correction for h in row_heights_final]
  else:
    row_heights_final = row_heights

  # Logo LPEE centré dans le bloc "Laboratoire de Contrôle Externe"
  logo_path = trouver_logo_lpee()
  if logo_path:
    try:
      from reportlab.platypus import Image as RLImage

      larg_bloc = sum(col_widths[:4])
      hauts = [h for h in row_heights_final[row1:row2 + 1] if h]
      haut_bloc = sum(hauts) if len(hauts) == 2 else 34
      iw, ih = ImageReader(logo_path).getSize()
      logo_h = max(haut_bloc - 6, 10)
      logo_w = logo_h * iw / ih
      if logo_w > larg_bloc * 0.4:
        logo_w = larg_bloc * 0.4
        logo_h = logo_w * ih / iw
      img = RLImage(logo_path, width=logo_w, height=logo_h)
      col_logo = logo_w + 10
      cellule = Table(
          [[img, P("Laboratoire de Contrôle Externe", size=9, bold=True, color=WHITE)]],
          colWidths=[col_logo, larg_bloc - col_logo],
          rowHeights=[haut_bloc],
      )
      cellule.setStyle(TableStyle([
          ("BACKGROUND", (0, 0), (-1, -1), DARK),
          ("ALIGN", (0, 0), (0, 0), "LEFT"),
          ("ALIGN", (1, 0), (1, 0), "CENTER"),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("LEFTPADDING", (0, 0), (0, 0), 5),
          ("LEFTPADDING", (1, 0), (1, 0), 0),
          ("RIGHTPADDING", (0, 0), (-1, -1), 0),
          ("TOPPADDING", (0, 0), (-1, -1), 0),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
          ("GRID", (0, 0), (-1, -1), 0, DARK),
      ]))
      data[row1][0] = cellule
    except Exception:
      pass  # le PV reste généré même si le logo pose problème

  table = Table(data, colWidths=col_widths, rowHeights=row_heights_final)
  table.setStyle(table_style)

  doc.build([table])
  buf.seek(0)
  return buf.getvalue()


def exporter_dataframe_excel(df, date_chaine):
  buf = io.BytesIO()
  with pd.ExcelWriter(buf, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name=f"Planning_{date_chaine}"[:31])
  buf.seek(0)
  return buf


# ==============================================================================
# 3. HELPER SUPABASE
# ==============================================================================
def obtenir_historique_betonnage(supabase, betonnage_id):
  if not betonnage_id:
    return []
  try:
    res = (
        projets_config.filtrer_projet_actif(
            supabase.table("suivi_controle_beton")
            .select("*")
            .eq("betonnage_id", betonnage_id)
        ).order("id").execute()
    )
    return res.data or []
  except Exception:
    return []


def obtenir_infos_betonnage_parent(supabase, betonnage_id):
  if not betonnage_id:
    return {}
  try:
    res = (
        projets_config.filtrer_projet_actif(
            supabase.table("suivi_betonnage")
            .select("*")
            .eq("id", betonnage_id)
        ).execute()
    )
    return res.data[0] if res.data else {}
  except Exception:
    return {}


def obtenir_infos_betonnage_parents_bulk(supabase, betonnage_ids):
  ids_valides = sorted({int(b) for b in betonnage_ids if pd.notnull(b)})
  if not ids_valides:
    return {}
  try:
    res = (
        projets_config.filtrer_projet_actif(
            supabase.table("suivi_betonnage")
            .select("*")
            .in_("id", ids_valides)
        ).execute()
    )
    return {p["id"]: p for p in (res.data or [])}
  except Exception as e:
    st.warning(f"Note : chargement groupé des fiches parentes impossible ({e}).")
    return {}


def determiner_ref_controle(supabase, betonnage_id, info_betonnage, sample_ep):
  key = f"ref_controle_beton_{betonnage_id}"

  num_rec = (info_betonnage or {}).get("num_reception")
  if num_rec and str(num_rec).strip().upper() not in [
      "",
      "-",
      "NONE",
      "NAN",
      "N/A",
  ]:
    ref = str(num_rec).strip()
    st.session_state[key] = ref
    return ref

  candidat = (info_betonnage or {}).get("ref_controle") or (sample_ep or {}).get("ref_controle")
  if candidat and str(candidat).strip():
    ref = str(candidat).strip()
    st.session_state[key] = ref
    return ref

  if key in st.session_state and st.session_state[key]:
    return st.session_state[key]

  ref = f"REF-{betonnage_id}-{(info_betonnage or {}).get('ouvrage', 'N/A')}"
  st.session_state[key] = ref
  return ref


# ==============================================================================
# 4. APPLICATION STREAMLIT
# ==============================================================================
def show(supabase):
  st.title("📋 Historique & Procès-Verbaux d'Écrasement (NF EN 12390)")

  user_info = st.session_state.get("user", {})
  role = str(
      st.session_state.get("user_role")
      or st.session_state.get("role")
      or user_info.get("role", "")
  ).lower()
  can_edit = st.session_state.get("can_edit", False) or bool(
      user_info.get("can_edit", False)
  )
  is_admin = (
      role in ["admin", "responsable_labo"]
      or st.session_state.get("is_admin", False)
  )
  is_baallal_admin = str(user_info.get("username", "")).strip().upper() == "BAALLAL" and role == "admin"

  if (
      role not in ["laboratoire", "coordinateur_essais", "labo", "admin", "responsable_labo", "responsable_chantier", "qualite"]
      and not is_admin
      and not can_edit
  ):
    st.error("⛔ **Accès Restreint**")
    st.warning(
        "Ce module est réservé exclusivement au personnel du **Laboratoire de"
        " Contrôle**."
    )
    return

  st.subheader("📋 Historique Général & Consultation des PVs")

  projet_id_actif = projets_config.projet_actif(user_info)
  if not projet_id_actif:
    st.error("⚠️ Aucun projet ne vous est autorisé. Contactez un administrateur.")
    return
  st.caption(f"📁 Projet actif : **{projets_config.nom_projet(projet_id_actif)}**")

  try:
    res_all = (
        supabase.table("suivi_controle_beton")
        .select("*")
        .eq("projet_id", projet_id_actif)
        .order("id", desc=True)
        .execute()
    )
    if not res_all.data:
      st.info("ℹ️ Aucun historique disponible dans la base de données.")
      return

    df_all = pd.DataFrame(res_all.data)

    unique_b_ids = [
        b_id for b_id in df_all["betonnage_id"].unique() if pd.notnull(b_id)
    ]
    unique_parents = obtenir_infos_betonnage_parents_bulk(supabase, unique_b_ids)

    def est_valide_val(v):
      if isinstance(v, bool):
        return v
      if isinstance(v, str):
        v_norm = (
            unicodedata.normalize("NFKD", v.strip().lower())
            .encode("ascii", "ignore")
            .decode("ascii")
        ).strip()
        if v_norm in ["true", "1", "ok", "oui"]:
          return True
        if any(mot in v_norm for mot in ["invalide", "non valide", "rejet"]):
          return False
        return "valide" in v_norm
      return False

    def verifier_pv_valide_et_signe(row):
      b_id = row.get("betonnage_id")
      f_kn = row.get("force_kn")
      try:
        a_force = pd.notnull(f_kn) and float(f_kn) > 0
      except (ValueError, TypeError):
        a_force = False

      # La validation est portée PAR LOT (bétonnage + échéance), donc sur
      # chaque éprouvette, et non plus sur la fiche parente : un lot validé
      # à 7 jours ne valide pas automatiquement les éprouvettes à 28 jours.
      statut_valide = est_valide_val(row.get("statut_pv"))

      return a_force and statut_valide

    mask_valides = df_all.apply(verifier_pv_valide_et_signe, axis=1)
    df_valides = df_all[mask_valides].copy()

    st.markdown("##### 📥 Re-télécharger un Procès-Verbal")

    cles_dans_liste = (
        set(
            zip(
                df_valides["betonnage_id"],
                df_valides["echeance"].map(lambda e: str(e).strip()),
            )
        )
        if not df_valides.empty
        else set()
    )
    lots_manquants = []
    groupes_lot_echeance = {}
    for _, r in df_all.iterrows():
      cle_groupe = (r.get("betonnage_id"), str(r.get("echeance", "-")).strip())
      groupes_lot_echeance.setdefault(cle_groupe, []).append(r)

    for (b_id, echeance_grp), rows_grp in groupes_lot_echeance.items():
      statut_lot = next(
          (r.get("statut_pv") for r in rows_grp if est_valide_val(r.get("statut_pv"))),
          None,
      )
      statut_admin_valide = est_valide_val(statut_lot)
      if statut_admin_valide and (b_id, echeance_grp) not in cles_dans_liste:
        a_au_moins_une_force = any(
            pd.notnull(r.get("force_kn")) and float(r.get("force_kn") or 0) > 0
            for r in rows_grp
        )
        lots_manquants.append({
            "Lot ID": b_id,
            "Échéance": echeance_grp,
            "Statut (admin)": statut_lot,
            "Au moins 1 force > 0 ?": "Oui" if a_au_moins_une_force else "Non",
        })

    if lots_manquants:
      with st.expander(
          f"🔧 {len(lots_manquants)} PV marqué(s) validé(s) mais absent(s) de"
          " la liste ci-dessous",
          expanded=True,
      ):
        st.caption(
            "Un PV validé n'apparaît dans la liste de téléchargement que si"
            " au moins une éprouvette de ce lot a une **Force (kN) > 0**. Ces"
            " lots sont marqués validés côté admin mais aucune éprouvette du"
            " lot n'a encore de force enregistrée :"
        )
        st.dataframe(
            pd.DataFrame(lots_manquants),
            use_container_width=True,
            hide_index=True,
        )

    if df_valides.empty:
      st.info(
          "ℹ️ Aucun Procès-Verbal **validé** n'est disponible pour le"
          " téléchargement."
      )
    else:
      c_r1, c_r2 = st.columns(2)
      recherche_pv = c_r1.text_input(
          "🔍 Rechercher (réf, ouvrage, classe...)",
          placeholder="Ex: gare casa sud, B/394...",
          key="search_input_pv",
      )
      recherche_date_pv = c_r2.text_input(
          "📅 Rechercher par Date d'écrasement",
          placeholder="Ex: 2026-08-08",
          key="search_date_pv",
      )

      groupes_valides = {}
      for _, row in df_valides.iterrows():
        b_id = row.get("betonnage_id")
        info_b = unique_parents.get(b_id) or {}
        ref_ctrl = determiner_ref_controle(
            supabase, b_id, info_b, row.to_dict()
        )
        classe = (
            row.get("classe_beton")
            or (info_b.get("classe_beton") or info_b.get("classe") if info_b else "-")
            or "-"
        )

        cle = (
            f"Référence : {ref_ctrl} | Classe : {classe} | Ouvrage :"
            f" {row.get('ouvrage', '-')} | Échéance : {row.get('echeance', '28 jours')}"
            f" (Date : {row.get('date_ecrasement', '-')}) | Lot ID #{b_id}"
        )
        groupes_valides.setdefault(cle, []).append(row.to_dict())

      pvs_filtrés = [
          k
          for k in groupes_valides.keys()
          if (not recherche_pv or recherche_pv.lower() in k.lower())
          and (not recherche_date_pv or recherche_date_pv in k)
      ]

      if not pvs_filtrés:
        st.warning("Aucun PV validé ne correspond à votre recherche.")
      else:
        choix_pv = st.selectbox(
            "Sélectionnez le PV à consulter :",
            pvs_filtrés,
            key="select_pv_hist",
        )
        lot_hist = groupes_valides[choix_pv]
        sample_h = lot_hist[0]
        b_id_h = sample_h.get("betonnage_id")

        info_b_h = unique_parents.get(b_id_h) or {}
        # Le PV téléchargé couvre tout l'historique du bétonnage (7 j + 28 j) ;
        # seules les échéances validées sont proposées à la sélection.
        essais_h = obtenir_historique_betonnage(supabase, b_id_h) or lot_hist
        # Un résultat écrasé (force > 0) n'entre dans le PV que si SON lot
        # (échéance) a été validé : un PV validé à 7 jours ne doit pas
        # publier des résultats à 28 jours non encore validés.
        essais_h = [
            it for it in essais_h
            if float(it.get("force_kn") or 0.0) == 0.0
            or est_valide_val(it.get("statut_pv"))
        ]

        date_coulee_h = info_b_h.get("date_coulee") or sample_h.get("date_coulee")

        export_data_h = []
        for item in essais_h:
          sec = float(item.get("section") or 176.71)
          if sec > 1000:
            sec = sec / 100.0
          f_kn = float(item.get("force_kn") or 0.0)
          type_essai_item = str(item.get("type_essai") or "Compression").strip()
          if "fendage" in type_essai_item.lower() or type_essai_item.lower().startswith("traction"):
            m_dim = re.search(r"(\d+)\s*x\s*(\d+)", str(item.get("forme") or ""))
            diam_mm, long_mm = (float(m_dim.group(1)), float(m_dim.group(2))) if m_dim else (150.0, 300.0)
            fc_calc = round((2.0 * f_kn * 1000.0) / (3.14159265 * diam_mm * long_mm), 2) if f_kn > 0 else 0.0
          else:
            fc_calc = round((f_kn * 10.0) / sec, 1) if f_kn > 0 else 0.0
          fc = float(item.get("fc_mpa") or fc_calc)
          ref_p = str(item.get("ref_controle") or "").strip()
          rep_s = str(item.get("repere_eprouvette", f"/{item['id']}")).strip()
          dt_essai_item = item.get("date_ecrasement", "-")

          age_real = calculer_age_jours(
              date_coulee_h, dt_essai_item, item.get("age")
          )

          export_data_h.append({
              "repere_eprouvette": f"{ref_p}{rep_s}" if ref_p else rep_s,
              "forme": item.get("forme", "Cylindrique 150x300"),
              "type_essai": type_essai_item,
              "section": sec,
              "force_kn": f_kn,
              "fc_mpa": fc,
              "date_essai": dt_essai_item,
              "age": age_real,
              "statut": "En cours" if f_kn == 0 else "Réalisé",
          })

        num_bl_h = extraire_num_bl(sample_h, info_b_h, choix_pv)
        ouv_h = info_b_h.get("ouvrage") or sample_h.get("ouvrage")
        ref_ctrl_h = determiner_ref_controle(
            supabase, b_id_h, info_b_h, sample_h
        )

        infos_header_h = {
            "re_num": "25/260/LGV/ B/",
            "dossier": projets_config.dossier_pv("2025-260-05985-2025-0247"),
            "client": projets_config.client_projet(),
            "num_reception": ref_ctrl_h,
            "ref_controle": ref_ctrl_h,
            "num_bl": num_bl_h,
            "ouvrage": ouv_h,
            "lieu_prelevement": ouv_h,
            "classe_beton": sample_h.get("classe_beton", "C35/45"),
            "date_coulee": date_coulee_h,
            "affaissement": (
                info_b_h.get("affaissement") or info_b_h.get("slump")
            ),
            "temperature": (
                info_b_h.get("temperature") or info_b_h.get("temp_beton")
            ),
            "forme": sample_h.get("forme", "Cylindrique 150x300"),
            "centrale": (
                info_b_h.get("centrale")
                or info_b_h.get("centrale_beton")
                or sample_h.get("centrale")
            ),
            "observations": (
                sample_h.get("observations_admin")
                or info_b_h.get("observations_admin")
                or sample_h.get("observations")
            ),
            "technicien_prelevement": (
                info_b_h.get("technicien_prelevement")
                or info_b_h.get("preleve_par")
                or info_b_h.get("technicien")
                or sample_h.get("technicien")
            ),
        }

        nom_rec_clean = nettoyer_nom_fichier(ref_ctrl_h)
        date_fab_clean = formater_date_nom_fichier(date_coulee_h)
        nom_fichier_pv = f"PV_{nom_rec_clean}_{date_fab_clean}.pdf"

        st.download_button(
            label=f"📄 Télécharger le PV ({nom_fichier_pv})",
            data=generer_pv_pdf(export_data_h, infos_header_h),
            file_name=nom_fichier_pv,
            mime="application/pdf",
            use_container_width=True,
            type="primary",
            key="btn_download_hist",
        )

        if is_baallal_admin:
          nom_fichier_pv_xlsx = f"PV_{nom_rec_clean}_{date_fab_clean}.xlsx"
          st.download_button(
              label=f"📊 Télécharger le PV en Excel ({nom_fichier_pv_xlsx}) — BAALLAL",
              data=generer_pv_excel(export_data_h, infos_header_h),
              file_name=nom_fichier_pv_xlsx,
              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              use_container_width=True,
              key="btn_download_hist_excel",
          )

    st.markdown("---")
    st.markdown("##### 📊 Base de données globale")

    df_all["ref_controle"] = df_all.apply(
        lambda r: determiner_ref_controle(
            supabase,
            r.get("betonnage_id"),
            unique_parents.get(r.get("betonnage_id")),
            r.to_dict(),
        ),
        axis=1,
    )

    df_all["affaissement_mm"] = df_all["betonnage_id"].map(
        lambda b: (unique_parents.get(b) or {}).get("affaissement")
        or (unique_parents.get(b) or {}).get("slump")
        or "-"
    )
    df_all["temp_beton_C"] = df_all["betonnage_id"].map(
        lambda b: (unique_parents.get(b) or {}).get("temperature")
        or (unique_parents.get(b) or {}).get("temp_beton")
        or "-"
    )

    df_all["statut_validation"] = df_all.apply(
        lambda r: (
            "✅ Validé & Signé"
            if verifier_pv_valide_et_signe(r)
            else "⏳ En attente"
        ),
        axis=1,
    )

    cols_ordre = [
        "id",
        "betonnage_id",
        "ref_controle",
        "repere_eprouvette",
        "num_bl",
        "ouvrage",
        "classe_beton",
        "statut_validation",
        "date_coulee",
        "affaissement_mm",
        "temp_beton_C",
        "echeance",
        "date_ecrasement",
        "fc_mpa",
        "technicien",
    ]
    exclus = {
        "forme",
        "section",
        "force_kn",
        "observations",
        "masse",
        "reference_controle",
        "refernce_controle",
        "num_reception",
    }

    cols_finales = [
        c
        for c in cols_ordre
        + [c for c in df_all.columns if c not in cols_ordre]
        if c not in exclus
    ]
    df_final = df_all[cols_finales]

    c_s1, c_s2 = st.columns(2)
    search_ref = c_s1.text_input(
        "🔍 Recherche par Réf. Contrôle",
        placeholder="Ex: REF-123-GARE CASA SUD",
    )
    search_date = c_s2.text_input(
        "📅 Recherche par Date de coulée", placeholder="Ex: 2026-08-24"
    )

    if search_ref:
      df_final = df_final[
          df_final["ref_controle"]
          .astype(str)
          .str.contains(search_ref, case=False, na=False)
      ]
    if search_date:
      df_final = df_final[
          df_final["date_coulee"]
          .astype(str)
          .str.contains(search_date, case=False, na=False)
      ]

    st.dataframe(df_final, use_container_width=True, hide_index=True)

    st.download_button(
        label="📊 Télécharger la base de données globale (Excel)",
        data=exporter_dataframe_excel(df_final, "Historique_Global"),
        file_name=f"Historique_Global_Beton_{date.today()}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="btn_download_hist_global",
    )

  except Exception as e:
    st.error(f"Erreur lors du chargement de l'historique global : {e}")
